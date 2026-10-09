"""Export every 2-, 4-, and 6-GPU H2D subset as per-GPU measurements."""

import csv
import json
import statistics
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1] / 'experiments/pcie_all8_concurrency_20261009'
PAYLOADS = ('Qwen expert', 'DeepSeek expert')
COUNTS = (2, 4, 6)
HEADERS = (
    ['Active GPUs', 'GPUs from 0–3']
    + [f'GPU{g} GiB/s' for g in range(8)]
    + [f'GPU{g} drop %' for g in range(8)]
    + ['Slowest GPU GiB/s', 'Fastest GPU GiB/s', 'GPU spread GiB/s',
       'Largest drop %', 'Aggregate repeat difference %']
)


def main():
    result = json.loads((ROOT / 'RESULTS.json').read_text())
    raw = json.loads((ROOT / 'RAW.json').read_text())
    assert result['status'] == 'PASS' and result['physical_gpus'] == list(range(8))
    assert len(raw) == 1020 and result['count'] == 510
    by = {(tuple(r['gpus']), r['payload']): r for r in result['rows']}
    repeats = {}
    for row in raw:
        repeats.setdefault((tuple(row['gpus']), row['payload']), []).append(row)
    solo = {(g, payload): by[((g,), payload)]['rank_gib_per_s'][str(g)]['median']
            for g in range(8) for payload in PAYLOADS}

    book = Workbook()
    intro = book.active
    intro.title = 'Read me'
    intro.append(['All 2-, 4-, and 6-GPU pinned-host H2D combinations'])
    intro.append(['Source', 'RESULTS.json and RAW.json from the completed all-eight study'])
    intro.append(['Cell definition', 'CUDA-event rank-local GiB/s; median of two repeats'])
    intro.append(['Drop definition', '100 × (1 − subset GPU speed / same-GPU solo speed)'])
    intro.append(['Negative drop', 'A small measured improvement versus solo; retain as measured'])
    intro.append(['Inactive GPU', 'Blank cell'])
    intro.append(['Aggregate repeat difference',
                  '100 × |repeat 1 − repeat 2| / mean(repeat 1, repeat 2)'])
    intro.append(['Scope', 'H2D copy only; no model inference, expert compute, NCCL, or prefetch'])
    intro.column_dimensions['A'].width = 31
    intro.column_dimensions['B'].width = 85
    intro['A1'].font = Font(bold=True, size=14)
    for row in intro.iter_rows(min_row=2, max_col=1):
        row[0].font = Font(bold=True)

    output = []
    for count in COUNTS:
        for payload in PAYLOADS:
            rows = [r for r in result['rows']
                    if len(r['gpus']) == count and r['payload'] == payload]
            assert len(rows) == {2: 28, 4: 70, 6: 28}[count]
            rows.sort(key=lambda r: (-sum(g < 4 for g in r['gpus']), r['gpus']))
            sheet = book.create_sheet(f'{count}GPU {"Qwen" if payload == PAYLOADS[0] else "DeepSeek"}')
            sheet.append(HEADERS)
            sheet.freeze_panes = 'C2'
            sheet.auto_filter.ref = f'A1:{get_column_letter(len(HEADERS))}{len(rows)+1}'
            sheet.row_dimensions[1].height = 31
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='23354D')
                cell.alignment = Alignment(wrap_text=True, vertical='center')
            sheet.column_dimensions['A'].width = 23
            sheet.column_dimensions['B'].width = 16
            for col in range(3, len(HEADERS) + 1):
                sheet.column_dimensions[get_column_letter(col)].width = 17
            for row in rows:
                group = tuple(row['gpus'])
                samples = repeats[group, payload]
                assert len(samples) == 2
                values = [sample['aggregate_gib_per_s'] for sample in samples]
                repeat_difference = 100 * abs(values[0] - values[1]) / statistics.mean(values)
                speed = {g: row['rank_gib_per_s'][str(g)]['median'] for g in group}
                drop = {g: 100 * (1 - speed[g] / solo[g, payload]) for g in group}
                output_row = (
                    ['·'.join(map(str, group)), sum(g < 4 for g in group)]
                    + [speed.get(g) for g in range(8)]
                    + [drop.get(g) for g in range(8)]
                    + [min(speed.values()), max(speed.values()),
                       max(speed.values()) - min(speed.values()), max(drop.values()),
                       repeat_difference]
                )
                sheet.append(output_row)
                output.append(dict(payload=payload, simultaneous_gpus=count,
                                   **dict(zip(HEADERS, output_row))))
            for cell_row in sheet.iter_rows(min_row=2, min_col=3):
                for cell in cell_row:
                    cell.number_format = '0.00'
            sheet.conditional_formatting.add(
                f'K2:R{len(rows)+1}',
                ColorScaleRule(start_type='num', start_value=0, start_color='E2F0D9',
                               mid_type='num', mid_value=10, mid_color='FFE699',
                               end_type='num', end_value=32, end_color='F4B6B6'))

    assert len(output) == 2 * (28 + 70 + 28)
    csv_path = ROOT / 'PER_RANK_2_4_6.csv'
    with csv_path.open('w', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=['payload', 'simultaneous_gpus'] + HEADERS,
                                lineterminator='\n')
        writer.writeheader()
        writer.writerows(output)
    book.save(ROOT / 'PER_RANK_2_4_6.xlsx')
    lines = [
        '# Per-GPU effective H2D for 2, 4, and 6 simultaneous GPUs',
        '',
        'Every one of the 28 two-GPU, 70 four-GPU, and 28 six-GPU subsets is '
        'included at both expert payload sizes in `PER_RANK_2_4_6.xlsx` and '
        '`PER_RANK_2_4_6.csv`. Blank GPU cells mean that GPU was inactive.',
        '',
        'The rates below are medians of CUDA-event rank-local GiB/s across '
        'the subsets in each row. Percentages are median losses against that '
        'same physical GPU measured alone at the same payload size. Each '
        'subset has two unfiltered repeats.',
        '',
        '| Simultaneous GPUs | From GPUs 0–3 | Subsets | '
        'Qwen 0–3 GiB/s (loss) | Qwen 4–7 GiB/s (loss) | '
        'DeepSeek 0–3 GiB/s (loss) | DeepSeek 4–7 GiB/s (loss) |',
        '|---:|---:|---:|---:|---:|---:|---:|',
    ]
    for count in COUNTS:
        for first_group in range(max(0, count - 4), min(4, count) + 1):
            cells = []
            subset_count = None
            for payload in PAYLOADS:
                subsets = [r for r in result['rows']
                           if len(r['gpus']) == count and r['payload'] == payload
                           and sum(g < 4 for g in r['gpus']) == first_group]
                subset_count = len(subsets)
                for front in (True, False):
                    pairs = [(r['rank_gib_per_s'][str(g)]['median'],
                              100 * (1 - r['rank_gib_per_s'][str(g)]['median'] / solo[g, payload]))
                             for r in subsets for g in r['gpus'] if (g < 4) == front]
                    cells.append(f'{statistics.median(v for v, _ in pairs):.2f} '
                                 f'({statistics.median(d for _, d in pairs):.1f}%)'
                                 if pairs else '—')
            lines.append(f'| {count} | {first_group} | {subset_count} | '
                         + ' | '.join(cells) + ' |')
    lines += [
        '',
        'Loss is 100 × (1 − subset speed / solo speed); negative values in '
        'the full table are small measured improvements and are preserved.',
        'This is repeated pinned-host H2D copy service, not model inference '
        'speed or an isolated PCIe component measurement.',
    ]
    (ROOT / 'PER_RANK_2_4_6.md').write_text('\n'.join(lines) + '\n')
    print(f'Exported {len(output)} subset/payload rows to {csv_path.name} and Excel')


if __name__ == '__main__':
    main()
